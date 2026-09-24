"""Build a single-column, ATS-safe .docx resume from a JSON data file.

Deliberately avoids everything that trips up real ATS parsers: no tables,
no text boxes, no columns, no header/footer content, no images/icons. Just
a standard font, standard heading styles, and plain paragraphs/bullets in
document order -- see this skill's SKILL.md for the reasoning.

Usage:
    python build_ats_resume.py --data resume_data.json --out resume.docx

resume_data.json shape -- see SKILL.md for a full annotated example:
{
  "name": str,                       (required)
  "contact": {                       (optional; all sub-fields optional)
    "email": str, "phone": str, "location": str,
    "linkedin": str, "portfolio": str
  },
  "summary": str,                    (optional)
  "skills": [str, ...],              (optional)
  "experience": [                    (required, at least one entry)
    {
      "title": str, "company": str, "location": str,
      "start": str, "end": str, "bullets": [str, ...]
    }
  ],
  "education": [                     (optional)
    {
      "degree": str, "school": str, "location": str,
      "start": str, "end": str, "details": str (optional)
    }
  ],
  "certifications": [str, ...]       (optional)
}
"""

import argparse
import json
import sys

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)
NAME_SIZE = Pt(20)


def _set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = BODY_SIZE


def _add_heading(doc: Document, text: str) -> None:
    # level=1 (real Heading 1 style) so ATS parsers that key off semantic
    # headings -- not just bold text -- can identify the section.
    heading = doc.add_heading(text, level=1)
    for run in heading.runs:
        run.font.name = BODY_FONT
        run.font.color.rgb = None  # keep default (black/theme), no styling tricks


def _add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def build(data: dict) -> Document:
    doc = Document()
    _set_base_style(doc)

    # Name -- plain bold paragraph, NOT the "Title" style (some Title
    # styles render as an image-like text box in certain templates/
    # themes; a bold Normal-style paragraph is unambiguous plain text).
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(data["name"])
    name_run.bold = True
    name_run.font.size = NAME_SIZE
    name_run.font.name = BODY_FONT

    # Contact info -- in the BODY, never header/footer (many ATS parsers
    # skip header/footer content entirely).
    contact = data.get("contact") or {}
    contact_parts = [
        contact.get(k)
        for k in ("email", "phone", "location", "linkedin", "portfolio")
        if contact.get(k)
    ]
    if contact_parts:
        doc.add_paragraph(" | ".join(contact_parts))

    if data.get("summary"):
        _add_heading(doc, "Summary")
        doc.add_paragraph(data["summary"])

    if data.get("skills"):
        _add_heading(doc, "Skills")
        # A single comma-separated paragraph, not a table -- keeps this
        # section trivially parseable regardless of how many skills there are.
        doc.add_paragraph(", ".join(data["skills"]))

    experience = data.get("experience") or []
    if experience:
        _add_heading(doc, "Experience")
        for job in experience:
            line = doc.add_paragraph()
            title_run = line.add_run(f"{job.get('title', '')} — {job.get('company', '')}")
            title_run.bold = True
            date_range = f"{job.get('start', '')} - {job.get('end', '')}".strip(" -")
            location = job.get("location") or ""
            tail_bits = [b for b in (location, date_range) if b]
            if tail_bits:
                line.add_run("  (" + ", ".join(tail_bits) + ")")
            for bullet in job.get("bullets") or []:
                _add_bullet(doc, bullet)

    education = data.get("education") or []
    if education:
        _add_heading(doc, "Education")
        for edu in education:
            line = doc.add_paragraph()
            degree_run = line.add_run(f"{edu.get('degree', '')} — {edu.get('school', '')}")
            degree_run.bold = True
            date_range = f"{edu.get('start', '')} - {edu.get('end', '')}".strip(" -")
            location = edu.get("location") or ""
            tail_bits = [b for b in (location, date_range) if b]
            if tail_bits:
                line.add_run("  (" + ", ".join(tail_bits) + ")")
            if edu.get("details"):
                doc.add_paragraph(edu["details"])

    if data.get("certifications"):
        _add_heading(doc, "Certifications")
        for cert in data["certifications"]:
            _add_bullet(doc, cert)

    return doc


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", required=True, help="Path to resume_data.json")
    parser.add_argument("--out", required=True, help="Output .docx path")
    args = parser.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    if not data.get("name"):
        sys.exit("resume_data.json must include a top-level \"name\".")
    if not data.get("experience"):
        sys.exit("resume_data.json must include at least one \"experience\" entry.")

    doc = build(data)
    doc.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
